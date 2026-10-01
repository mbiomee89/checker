import openpyxl
VAT=0.15; COM=0.21; PAY=0.025; PCT=COM+PAY; HP=4.0
K=1/(1+VAT)-PCT   # economic net per 1 SAR of customer-paid price, before fixed delivery contribution
BANDS=[(30,40,7),(40,50,8),(50,60,9),(60,80,11),(80,1e9,13)]
def D(P, under30=7):
    if P<30: return under30
    for a,b,d in BANDS:
        if a<=P<b: return d
def net_del(P,h=0,under30=7): return P*K-D(P,under30)-HP*h
def net_pick(P): return P*K-min(2.5,0.2*P)   # discount share borne by restaurant (see assumptions)
def direct_ex(p): return p/(1+VAT)
def be_price(p,h=0,under30=7,k=K):
    P=p/1.15/k
    for _ in range(50):
        P=(p/1.15+D(P,under30)+HP*h)/k
    return P
def be_pick(p): return (p/1.15+2.5)/K if p>12.5 else p/1.15/(K+0.0)  # approx for tiny
def load():
    wb=openpyxl.load_workbook('/root/.claude/uploads/c2965937-9e4a-591e-9cc6-a7dfaddcf750/70348073-Musa_2.xlsx',data_only=True)
    ws=wb.worksheets[0]; rows=[]
    for r in range(2,58):
        rows.append(dict(row=r,sec=ws.cell(r,1).value,name=ws.cell(r,2).value.strip(),
            D=(float(ws.cell(r,4).value) if ws.cell(r,4).value not in (None,'') else None),F=float(ws.cell(r,6).value)))
    return rows
MENU={2:80,3:77,4:75,5:36,6:45,7:14,8:21,9:42,10:None,11:21,12:21,13:15,14:15,15:42,16:42,17:21,18:18,19:50,20:27,21:25,22:10,23:None,24:10,25:15,26:7,27:7,28:8,29:10,30:17,31:8,32:45,33:17,34:40,35:20,36:35,37:35,38:35,39:45,40:15,41:15,42:10,43:8,44:6,45:5,46:2,47:8,48:5,49:5,50:4,51:5,52:5,53:2.5,54:2.5,55:2.5,56:2,57:1}
if __name__=='__main__':
    print('K',K)
    for p in (30,40,50,60,80,100,120,150,200):
        print(p,[round(be_price(p,h)/p-1,3) for h in (0,1)], 'pick',round(be_pick(p)/p-1,3))
