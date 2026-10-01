import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from model import *
rows=load()
wb=openpyxl.Workbook()
A=wb.active; A.title='Assumptions'
hdr=PatternFill('solid',fgColor='1F3864'); hf=Font(color='FFFFFF',bold=True); inp=PatternFill('solid',fgColor='FFF2CC')
A.sheet_view.rightToLeft=True
data=[('VAT','ضريبة القيمة المضافة',0.15,'مقرر نظاماً (ليس من العقد) - الأسعار مفترضة شاملة الضريبة'),
('COM','عمولة هنقرستيشن (توصيل/استلام)',0.21,'عقد ص4 البند خامساً-أ'),
('PAY','عمولة الدفع الإلكتروني',0.025,'عقد ص4 البند خامساً-أ'),
('ESH','نسبة الطلبات المدفوعة إلكترونياً',1.0,'افتراض متحفظ: 100% (غير محدد في العقد)'),
('HPLUS','إضافة HPlus لكل طلب (ريال)',4,'عقد ص6'),
('FOOD','تكلفة الطعام % من سعر البيع المباشر بدون ضريبة',0.32,'افتراض - يجب استبداله بتكلفة الوصفات'),
('PKG','تغليف لكل طلب (ريال)',2,'افتراض - يجب استبداله'),
('PMAX','سقف حصة المطعم من خصم الاستلام (ريال)',2.5,'عقد ص5'),
('PDISC','خصم العميل على الاستلام',0.20,'عقد ص5 ("يبدأ من 20%" ويحق لهنقرستيشن تغييره)'),
('U30','تكلفة التوصيل لسلة أقل من 30 ريال',7,'غير مذكورة بالعقد - افتراض (أقل شريحة)'),
('SUB','الاشتراك الشهري (ريال)',200,'عقد ص4: 200 من 01/10/2026 إلى 01/02/2027 (غامض: شهري أم إجمالي)'),
('PRN','طابعة V2s مرة واحدة (ريال)',500,'عقد ص4')]
A.append(['الرمز','البند','القيمة','المصدر / ملاحظة'])
for c in A[1]: c.fill=hdr; c.font=hf
for d in data: A.append(list(d))
for r in range(2,2+len(data)): A.cell(r,3).fill=inp
A.append([]); A.append(['','شرائح تكلفة التوصيل (حسب قيمة السلة شاملة الضريبة)']); 
A.append(['من (ريال)','تكلفة التوصيل (ريال)'])
BS=A.max_row+1
for a,d in [(30,7),(40,8),(50,9),(60,11),(80,13)]: A.append([a,d])
A.column_dimensions['A'].width=10; A.column_dimensions['B'].width=48; A.column_dimensions['C'].width=12; A.column_dimensions['D'].width=70
ref={d[0]:f'Assumptions!$C${i+2}' for i,d in enumerate(data)}
band=f'Assumptions!$A${BS}:$A${BS+4}'; bandv=f'Assumptions!$B${BS}:$B${BS+4}'
I=wb.create_sheet('Items'); I.sheet_view.rightToLeft=True
heads=['#','الصنف','القسم','سعر المطعم الحالي (Excel عمود D)','سعر هنقرستيشن المقترح (Excel عمود F)','نسبة الزيادة','ضريبة 15% المضمنة في السعر','عمولة 21%','رسوم الدفع 2.5%','تكلفة التوصيل (شريحة - طلب منفرد)','إجمالي خصومات المنصة','الصافي للمطعم قبل التكلفة (بعد الضريبة)','الصافي % من سعر المطعم','الصافي % من سعر HS','الصافي % من البيع المباشر (بدون ضريبة)','صافي طلب HPlus','أقصى تكلفة طعام+تغليف قبل الخسارة','تكلفة طعام+تغليف (افتراض)','هامش المساهمة (CM)','CM % من سعر HS بدون ضريبة','التصنيف','CM البيع المباشر','الفرق عن البيع المباشر','صافي طلب الاستلام (Pickup)','سعر المنيو المطبوع','سعر HS ÷ سعر المنيو - 1','ملاحظة']
I.append(heads)
for c in I[1]: c.fill=hdr; c.font=hf; c.alignment=Alignment(wrap_text=True,vertical='center')
I.row_dimensions[1].height=60
V,CM_,PY,ES,HP_,FD,PK,PM,PD,U3=[ref[k] for k in ('VAT','COM','PAY','ESH','HPLUS','FOOD','PKG','PMAX','PDISC','U30')]
for i,r in enumerate(rows):
    n=i+2; m=MENU[r['row']]
    note=''
    d=r['D']
    if d is None: d=m; note='لا يوجد سعر D في الإكسل (خانة فارغة) - استُخدم سعر المنيو 2.5 كسعر مباشر؛ سعر HS=4 مكتوب يدوياً'
    elif abs(r['F']/r['D']-1.2)>0.005: note='سعر HS مكتوب يدوياً ولا يساوي D×1.2'
    if m is None: note=(note+'; ' if note else '')+'الصنف غير موجود في المنيو المطبوع'
    I.append([i+1,r['name'],r['sec'] or '',d,r['F'],f'=E{n}/D{n}-1',f'=E{n}*{V}/(1+{V})',f'=E{n}*{CM_}',f'=E{n}*{PY}*{ES}',
      f'=IF(E{n}<30,{U3},LOOKUP(E{n},{band},{bandv}))',f'=H{n}+I{n}+J{n}',f'=E{n}-G{n}-K{n}',f'=L{n}/D{n}',f'=L{n}/E{n}',f'=L{n}/(D{n}/(1+{V}))',f'=L{n}-{HP_}',f'=L{n}',
      f'={FD}*D{n}/(1+{V})+{PK}',f'=L{n}-R{n}',f'=S{n}/(E{n}/(1+{V}))',f'=IF(T{n}<0.15,"RED",IF(T{n}<0.3,"YELLOW","GREEN"))',f'=D{n}/(1+{V})-R{n}',f'=S{n}-V{n}',
      f'=E{n}-G{n}-E{n}*{CM_}-E{n}*{PY}*{ES}-MIN({PM},{PD}*E{n})',m if m is not None else '', f'=IF(Y{n}="","",E{n}/Y{n}-1)',note])
last=len(rows)+1
for col in 'FMNO': 
    for r in range(2,last+1): I[f'{col}{r}'].number_format='0.0%'
for r in range(2,last+1):
    I[f'T{r}'].number_format='0.0%'; I[f'Z{r}'].number_format='0.0%'
    for col in 'GHIJKLPQRSVWX': I[f'{col}{r}'].number_format='0.00'
I.conditional_formatting.add(f'U2:U{last}',CellIsRule(operator='equal',formula=['"RED"'],fill=PatternFill('solid',bgColor='F8CBAD')))
I.conditional_formatting.add(f'U2:U{last}',CellIsRule(operator='equal',formula=['"YELLOW"'],fill=PatternFill('solid',bgColor='FFE699')))
I.conditional_formatting.add(f'U2:U{last}',CellIsRule(operator='equal',formula=['"GREEN"'],fill=PatternFill('solid',bgColor='C6E0B4')))
I.column_dimensions['B'].width=24; I.column_dimensions['AA'].width=60
I.freeze_panes='C2'
wb.save('model.xlsx')

# ---- append static analysis tables parsed from REPORT.md
import re
wb=openpyxl.load_workbook('model.xlsx')
S=wb.create_sheet('Analysis_Tables'); S.sheet_view.rightToLeft=True
def num(x):
    t=x.strip().replace(',','')
    try: return float(t)
    except: return x.strip()
for line in open('REPORT.md',encoding='utf8'):
    line=line.rstrip('\n')
    if line.startswith('#'):
        S.append([]); S.append([line.lstrip('# ').strip()]); S.cell(S.max_row,1).font=Font(bold=True,size=12)
    elif line.startswith('|') and not re.match(r'^\|[-| ]+\|$',line):
        cells=[num(c) for c in line.strip().strip('|').split(' | ')]
        S.append(cells)
S.column_dimensions['A'].width=30
wb.save('model.xlsx')
